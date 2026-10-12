"""Read-only source intake, per-SKU checkpoints, and verified delivery."""
from pathlib import Path
import argparse,datetime,hashlib,json,os,re,shutil,zipfile,xml.etree.ElementTree as ET
from PIL import Image

DEFAULT_DOC=Path(__file__).resolve().parents[1]/'references/design-reference-snapshot.docx'
DEFAULT_MODULES=[{'name':'campaign','kind':'campaign','width':1464,'height':1200,'people':3},{'name':'fit_detail','kind':'fit_detail','width':1464,'height':1200,'people':1},{'name':'styling','kind':'styling','width':1464,'height':600,'people':2,'looks':2},{'name':'creator_occasion','kind':'creator_occasion','width':1464,'height':600,'people':3,'solo_each':True}]
SKIP={'弃用','过程文件','流程准备','output','.git','.cache'}
def now():return datetime.datetime.now(datetime.timezone.utc).isoformat()
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def write(p,v):
 p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);t=p.with_suffix(p.suffix+'.tmp');t.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8');os.replace(t,p)
def sha(p):
 h=hashlib.sha256()
 with Path(p).open('rb') as f:
  for b in iter(lambda:f.read(4*1024*1024),b''):h.update(b)
 return h.hexdigest()
def slug(s):return re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',s).strip(' .')[:110] or 'sku'
def excluded(p):return any(x in SKIP or x.lower().startswith('aplus_') for x in p.parts)
def books(p):return sorted(p.glob('*产品属性表*.xlsx'))
def discover(sources):
 out=[]
 for raw in sources:
  p=Path(raw).resolve()
  if not p.is_dir():raise ValueError('Source directory unavailable: '+str(p))
  if books(p):out.append(p);continue
  for current,dirs,_ in os.walk(p):
   dirs[:]=[x for x in dirs if not excluded(Path(x))];q=Path(current)
   if books(q):out.append(q);dirs[:]=[]
 out=list(dict.fromkeys(out))
 if not out:raise ValueError('No SKU with a product attribute workbook was found')
 return out
def capture_rules(path,out,explicit=False):
 chosen=Path(path);fallback=False
 if not chosen.is_file():
  if explicit:raise ValueError('Specified reference document unavailable: '+str(chosen))
  chosen=Path(__file__).resolve().parents[1]/'references/design-reference-snapshot.docx';fallback=True
 digest=sha(chosen);out=Path(out)/digest[:16];out.mkdir(parents=True,exist_ok=True);shutil.copy2(chosen,out/'reference.docx');media=out/'case_images';media.mkdir(exist_ok=True)
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
 with zipfile.ZipFile(chosen) as z:
  rels={e.attrib['Id']:e.attrib['Target'] for e in ET.fromstring(z.read('word/_rels/document.xml.rels'))};doc=ET.fromstring(z.read('word/document.xml'));lines=[];blocks=[]
  for p in doc.findall('.//w:body//w:p',ns):
   text=''.join(t.text or '' for t in p.findall('.//w:t',ns));images=[rels.get(a.attrib.get('{'+ns['r']+'}embed'),'') for a in p.findall('.//a:blip',ns)]
   if text or images:blocks.append({'text':text,'images':images});lines.append(text+''.join(' [CASE '+x+']' for x in images))
  for n in z.namelist():
   if n.startswith('word/media/') and not n.endswith('/'):(media/Path(n).name).write_bytes(z.read(n))
 (out/'text.txt').write_text('\n'.join(lines),encoding='utf-8');write(out/'blocks.json',blocks)
 return {'requested_path':str(path),'actual_path':str(chosen),'captured_doc':str(out/'reference.docx'),'sha256':digest,'fallback_snapshot':fallback,'captured_utc':now(),'text':str(out/'text.txt'),'case_directory':str(media),'case_files':[p.name for p in media.iterdir()]}
def inventory(source):
 images=[]
 for current,dirs,files in os.walk(source):
  dirs[:]=[x for x in dirs if not excluded(Path(x))]
  for n in sorted(files):
   p=Path(current)/n
   if p.suffix.lower() not in {'.png','.jpg','.jpeg','.webp','.bmp','.tif','.tiff'}:continue
   with Image.open(p) as im:images.append({'relative_path':str(p.relative_to(source)),'absolute_path':str(p),'size':list(im.size),'actual_format':im.format,'mode':im.mode,'bytes':p.stat().st_size})
 return images
def attributes(path):
 import openpyxl
 # Normal-mode openpyxl discards non-anchor values in merged ranges.
 # Some WPS workbooks store real section labels/copy in those XML cells.
 wb=openpyxl.load_workbook(path,data_only=False,read_only=True)
 cache=openpyxl.load_workbook(path,data_only=True,read_only=True)
 metadata=openpyxl.load_workbook(path,data_only=False);out={}
 for ws in wb.worksheets:
  ws.reset_dimensions();cached=cache[ws.title];cached.reset_dimensions()
  cached_values={c.coordinate:c.value for row in cached for c in row if c.value is not None}
  info=metadata[ws.title];links={c.coordinate:c.hyperlink.target for row in info for c in row if c.hyperlink}
  cells=[];seen=set()
  for row in ws:
   for c in row:
    if c.value is not None:
     seen.add(c.coordinate);cells.append({'cell':c.coordinate,'value':c.value,'cached_value':cached_values.get(c.coordinate),'hyperlink':links.get(c.coordinate)})
  for coordinate,url in links.items():
   if coordinate not in seen:cells.append({'cell':coordinate,'value':None,'cached_value':cached_values.get(coordinate),'hyperlink':url})
  out[ws.title]={'range':info.calculate_dimension(),'cells':cells,'merged_ranges':[str(r) for r in info.merged_cells.ranges],'reader':'read_only_full_xml_cells'}
 wb.close();cache.close();metadata.close();return out
def validate_modules(modules):
 names=[m['name'] for m in modules]
 if not modules or len(names)!=len(set(names)):raise ValueError('Modules must have unique names')
 widths={m.get('width',1464) for m in modules}
 if len(widths)!=1:raise ValueError('A stacked PSD requires a consistent module width; resolve the user requirement before generation')
 for m in modules:
  if not isinstance(m.get('width',1464),int) or not 0<m.get('width',1464)<=30000 or not isinstance(m['height'],int) or m['height']<=0:raise ValueError('Modules require positive integer dimensions within PSD limits')
 return modules
def initialize(args):
 config=read(args.config) if args.config else {};sources=args.source or config.get('sources',[])
 if not sources:raise ValueError('Provide --source or config.sources')
 root=Path(args.work_root).resolve();root.mkdir(parents=True,exist_ok=True)
 if (root/'batch.json').exists():raise ValueError('Workspace already exists; inspect existing batch.json and resume instead of reinitializing')
 modules=validate_modules(read(args.modules) if args.modules else config.get('modules',DEFAULT_MODULES))
 rules=capture_rules(args.rules_doc or config.get('rules_doc',str(DEFAULT_DOC)),root/'rules',explicit=bool(args.rules_doc or config.get('rules_doc')))
 batch={'created_utc':now(),'rules':rules,'items':[],'status':'running'}
 for i,source in enumerate(discover(sources),1):
  ws=root/'skus'/f'{i:03d}_{slug(source.name)}';ws.mkdir(parents=True,exist_ok=True)
  for folder in ['assets','references','prompts','checks','scripts']:(ws/folder).mkdir(exist_ok=True)
  overrides=config.get('sku_overrides',{}).get(source.name,{});mods=validate_modules(overrides.get('modules',modules));tables=books(source)
  a=max([m.get('people',3 if m.get('kind')=='campaign' else 2 if m.get('kind')=='styling' else 1) for m in mods if m.get('kind') not in {'creator','occasion','creator_occasion'}]+[0]);b=max([m.get('people',3) for m in mods if m.get('kind') in {'creator','occasion','creator_occasion'}]+[0])
  m={'sku':source.name,'source_directory':str(source),'workspace':str(ws),'batch_file':str(root/'batch.json'),'rules':rules,'modules':mods,'canvas':{'width':mods[0].get('width',1464),'height':sum(x['height'] for x in mods)},'model_policy':{'A_count':a,'B_count':b,'disjoint':True,'creator_solo_each':True},'usage':config.get('usage','personal_practice'),'output_root':args.output_root or config.get('output_root'),'product_lock':{},'model_identities':[],'scene_research':[],'pairing_products':[],'jobs':{},'qc':{},'status':'input_ready','overrides':overrides}
  try:
   if len(tables)!=1:raise ValueError('Expected one product attribute workbook; found '+str(len(tables)))
   m['attribute_workbook']=str(tables[0]);m['attribute_workbook_sha256']=sha(tables[0]);data=attributes(tables[0]);write(ws/'attributes.json',data)
   from copy_contract import extract_candidates
   candidates=extract_candidates(data);write(ws/'copy_candidates.json',candidates)
   write(ws/'copy_plan.json',{'reviewed':False,'entries':[dict(id=c['id'],decision='pending',placements=[]) for c in candidates]})
   m['inventory']=inventory(source)
  except Exception as e:m.update(status='blocked',error=str(e))
  write(ws/'manifest.json',m);batch['items'].append({'sku':source.name,'workspace':str(ws),'status':m['status']});print(m['status'],source.name,str(ws))
 write(root/'batch.json',batch)
def load_workspace(path):
 ws=Path(path).resolve();return ws,read(ws/'manifest.json')
def checkpoint(ws,m):
 m['updated_utc']=now();write(ws/'manifest.json',m)
 if m.get('batch_file') and Path(m['batch_file']).exists():
  b=read(m['batch_file'])
  for item in b['items']:
   if item['workspace']==str(ws):item.update(status=m['status'],delivery_directory=m.get('delivery_directory'))
  b['status']='complete' if all(x['status']=='delivered' for x in b['items']) else 'running';write(m['batch_file'],b)
def record_job(args):
 ws,m=load_workspace(args.workspace);j=m['jobs'].setdefault(args.key,{'created_utc':now()})
 if j.get('status')=='completed' and args.status!='completed':raise ValueError('Completed job key is reserved; recover it or use a new versioned key')
 if j.get('status') in {'submitted','running','unknown'} and args.status=='planned':raise ValueError('Existing job may still be running; inspect recorded handle instead of planning a duplicate')
 if j.get('status')=='completed' and args.result and j.get('sha256')!=sha(args.result):raise ValueError('Use a new versioned key for a different generation result')
 j.update(status=args.status,updated_utc=now())
 if args.handle:j['handle']=args.handle
 if args.prompt_file:
  p=Path(args.prompt_file).resolve();j.update(prompt_file=str(p),prompt_sha256=sha(p))
 if args.references:j['references']=read(args.references)
 if args.expected:j['expected']=args.expected
 if args.result:
  source=Path(args.result).resolve();dst=ws/'assets'/slug(args.asset_name or source.name)
  if dst.exists() and sha(dst)!=sha(source):raise ValueError('Asset collision; provide a new versioned --asset-name')
  if source!=dst:shutil.copy2(source,dst)
  with Image.open(dst) as im:j.update(original_result=str(source),result=str(dst.relative_to(ws)),size=list(im.size),sha256=sha(dst))
  j['status']='completed'
 if j['status']=='completed' and not j.get('result'):raise ValueError('Completed job needs a validated saved result')
 m['status']='generating' if any(x['status'] in {'planned','submitted','running','unknown'} for x in m['jobs'].values()) else 'assets_ready' if all(x['status']=='completed' and x.get('result') for x in m['jobs'].values()) else 'generation_failed';checkpoint(ws,m);print(json.dumps(j,ensure_ascii=False))
def update(args):
 ws,m=load_workspace(args.workspace);patch=read(args.patch)
 if any(k in patch for k in ['source_directory','workspace','sku','rules']):raise ValueError('Source identity and captured criteria cannot be changed by a stage patch')
 m.update(patch)
 if 'modules' in patch:m['modules']=validate_modules(m['modules']);m['canvas'].update(width=m['modules'][0].get('width',1464),height=sum(x['height'] for x in m['modules']))
 checkpoint(ws,m);print(m['status'])
def refresh_rules(args):
 root=Path(args.work_root).resolve();b=read(root/'batch.json');old=b['rules'];path=Path(args.rules_doc or old['requested_path'])
 actual=path if path.is_file() else Path(__file__).resolve().parents[1]/'references/design-reference-snapshot.docx'
 explicit=bool(args.rules_doc or str(path)!=str(DEFAULT_DOC))
 if not path.is_file() and explicit:raise ValueError('Latest specified reference document unavailable')
 if sha(actual)==old['sha256'] and bool(not path.is_file())==old['fallback_snapshot']:print('CRITERIA_UNCHANGED');return
 new=capture_rules(path,root/'rules',explicit);b['rules']=new;write(root/'batch.json',b)
 for item in b['items']:
  ws,m=load_workspace(item['workspace'])
  if m['status']=='delivered':continue
  m.setdefault('criteria_history',[]).append(m['rules']);m.update(rules=new,criteria_read=False,status='criteria_changed');checkpoint(ws,m)
 print('CRITERIA_CHANGED — reread text and all cases before dependent work')
def deliver(args):
 ws,m=load_workspace(args.workspace);v=read(ws/'checks/validation.json')
 if v.get('status')!='passed':raise ValueError('Delivery validation not passed')
 for f,digest in v.get('artifact_sha256',{}).items():
  if not Path(f).is_file() or sha(f)!=digest:raise ValueError('Artifact changed after validation; inspect and revalidate')
 for f,digest in v.get('copy_evidence_sha256',{}).items():
  if not Path(f).is_file() or sha(f)!=digest:raise ValueError('Copy evidence changed after validation; inspect and revalidate')
 paths=list(m.get('outputs',read(ws/'checks/build_status.json').get('outputs',{})).values());paths=[ws/p for p in paths];notes=ws/'sources.md'
 if not notes.exists():raise ValueError('Create sources.md before delivery')
 paths.append(notes)
 if m.get('delivery_directory'):
  old=Path(m['delivery_directory']);receipt=read(old/'delivery_receipt.json')
  if not all((old/f['name']).is_file() and sha(old/f['name'])==f['sha256'] for f in receipt['files']):raise ValueError('Existing delivery changed outside this workflow; inspect it first')
  current={p.name:sha(p) for p in paths};previous={f['name']:f['sha256'] for f in receipt['files']}
  if current==previous:print('EXISTING_VERIFIED',str(old));return
 base=Path(m['output_root'])/slug(m['sku']) if m.get('output_root') else Path(m['source_directory']);base.mkdir(parents=True,exist_ok=True)
 stamp=datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ');dest=base/('Aplus_PSD_'+stamp)
 dest.mkdir(exist_ok=False);receipt={'files':[],'created_utc':now()}
 for src in paths:
  target=dest/src.name;shutil.copy2(src,target);digest=sha(src)
  if digest!=sha(target):raise ValueError('Copy checksum mismatch '+src.name)
  receipt['files'].append({'name':src.name,'sha256':digest,'bytes':target.stat().st_size})
 prep=dest/'流程准备';prep.mkdir();shutil.copy2(ws/'manifest.json',prep/'manifest.json');shutil.copy2(ws/'attributes.json',prep/'attributes.json')
 for name in ['copy_candidates.json','copy_plan.json','layout.json']:
  if (ws/name).is_file():shutil.copy2(ws/name,prep/name)
 for name in ['prompts','checks','scripts']:shutil.copytree(ws/name,prep/name,dirs_exist_ok=True)
 write(dest/'delivery_receipt.json',receipt);m.update(status='delivered',delivery_directory=str(dest));checkpoint(ws,m);shutil.copy2(ws/'manifest.json',prep/'manifest.json');print('DELIVERED',str(dest))
def main():
 p=argparse.ArgumentParser();sub=p.add_subparsers(dest='command',required=True)
 a=sub.add_parser('init');a.add_argument('--source',action='append');a.add_argument('--config');a.add_argument('--work-root',required=True);a.add_argument('--output-root');a.add_argument('--rules-doc');a.add_argument('--modules');a.set_defaults(fn=initialize)
 a=sub.add_parser('job');a.add_argument('--workspace',required=True);a.add_argument('--key',required=True);a.add_argument('--status',choices=['planned','submitted','running','completed','failed','unknown'],required=True);a.add_argument('--handle');a.add_argument('--prompt-file');a.add_argument('--references');a.add_argument('--expected');a.add_argument('--result');a.add_argument('--asset-name');a.set_defaults(fn=record_job)
 a=sub.add_parser('update');a.add_argument('--workspace',required=True);a.add_argument('--patch',required=True);a.set_defaults(fn=update)
 a=sub.add_parser('deliver');a.add_argument('--workspace',required=True);a.set_defaults(fn=deliver)
 a=sub.add_parser('refresh-rules');a.add_argument('--work-root',required=True);a.add_argument('--rules-doc');a.set_defaults(fn=refresh_rules)
 args=p.parse_args();args.fn(args)
if __name__=='__main__':main()
